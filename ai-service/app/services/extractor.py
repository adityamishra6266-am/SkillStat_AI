# ai-service/app/services/extractor.py
# Extracts plain text from PDF, DOCX, PPTX and TXT files.

import io
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation


class ExtractionError(Exception):
    pass


def extract_text(file_bytes: bytes, file_type: str) -> str:
    """Dispatch to the right extractor based on file_type ('pdf', 'docx', 'pptx', 'txt')."""
    file_type = (file_type or "").lower().strip()

    try:
        if file_type == "pdf":
            return _extract_pdf(file_bytes)
        elif file_type == "docx":
            return _extract_docx(file_bytes)
        elif file_type == "pptx":
            return _extract_pptx(file_bytes)
        elif file_type == "txt":
            return _extract_txt(file_bytes)
        else:
            raise ExtractionError(f"Unsupported file type: {file_type}")
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Failed to extract text from {file_type} file: {exc}") from exc


def _extract_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also pull text from tables, if any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts).strip()


def _extract_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line.strip())
    return "\n".join(parts).strip()


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode text file with common encodings.")
